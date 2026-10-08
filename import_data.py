"""Convert the supplied faculty workbook into the public thesis fields used by the app."""
from datetime import date, datetime
import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "FKOM - Teknik Informatika S1 Laporan Prasidang (PIN) dan Pelaksanaan Sidang (1).xlsx"
OUTPUT = ROOT / "data" / "theses.json"

ADVISER_NAMES = {
    "aahsumiah": "Aah Sumiah, M.Kom.",
    "aguswahyuddin": "Agus Wahyuddin, M.Kom.",
    "ajipermana": "Aji Permana, M.Kom.",
    "cecepjabbas": "Cecep J. Abbas, M.Kom.",
    "dadanghamdani": "Dadang Hamdani, M.Kom.",
    "erlandarmawan": "Erlan Darmawan, M.Si.",
    "fauziah": "Fauziah, M.Kom.",
    "fitranugraha": "Fitra Nugraha, M.Kom.",
    "herubudianto": "Heru Budianto, M.Kom.",
    "iwanlesmana": "Iwan Lesmana, M.Kom.",
    "jerrydounaldrahajaan": "Jerry Dounald Rahajaan, M.Sn.",
    "nidaamaliaasikin": "Nida Amalia Asikin, M.Pd.",
    "nununugraha": "Nunu Nugraha, M.T.",
    "panjinovantara": "Panji Novantara, M.T.",
    "rioandriyatkrisdiawan": "Rio Andriyat Krisdiawan, M.Kom.",
    "riopriantama": "Rio Priantama, M.T.I.",
    "sherlyginasupratman": "Sherly Gina Supratman, M.Kom.",
    "sitimaesyaroh": "Siti Maesyaroh, M.Kom.",
    "sugengsupriyadi": "Sugeng Supriyadi, M.Kom.",
    "titosugiharto": "Tito Sugiharto, M.Eng.",
    "yatinurhayati": "Yati Nurhayati, M.Kom.",
    "yulyanto": "Yulyanto, M.T.I.",
}
ADVISER_ALIASES = {
    "aguswahyudin": "aguswahyuddin",
    "dadanghamani": "dadanghamdani",
    "erlandharmawan": "erlandarmawan",
    "iwanlemana": "iwanlesmana",
    "jerrydounaldr": "jerrydounaldrahajaan",
    "rioandriyat": "rioandriyatkrisdiawan",
    "rioandriyatk": "rioandriyatkrisdiawan",
    "rioandryat": "rioandriyatkrisdiawan",
    "shelyginasupratman": "sherlyginasupratman",
    "titosugijarto": "titosugiharto",
}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def text(row, index):
    return clean(row[index]) if index < len(row) else ""


def adviser_key(value):
    name = re.split(r",|\.\s+(?=M\.)", value, maxsplit=1, flags=re.IGNORECASE)[0]
    return re.sub(r"[^a-z]", "", name.casefold())


def canonical_adviser(value):
    key = adviser_key(value)
    key = ADVISER_ALIASES.get(key, key)
    return ADVISER_NAMES.get(key, value.strip())


def year_of(value):
    if isinstance(value, (datetime, date)):
        return value.year
    raw = clean(value)
    for token in raw.replace("/", "-").split("-"):
        if token.isdigit() and len(token) == 4 and 2000 <= int(token) <= 2100:
            return int(token)
    return None


def main(source=SOURCE, output=OUTPUT):
    workbook = load_workbook(source, read_only=True, data_only=True)
    records = []
    for sheet in workbook.worksheets:
        legacy = sheet.title == "2020-2024"
        for row in sheet.iter_rows(min_row=4, values_only=True):
            name = text(row, 4 if legacy else 3)
            initial_title = text(row, 11 if legacy else 10)
            final_title = text(row, 27 if legacy else 26)
            title = final_title or initial_title
            if not title or not name:
                continue
            sidang = row[2] if legacy else row[1]
            graduation = row[23] if len(row) > 23 else None
            year = year_of(graduation) or year_of(sidang) or year_of(sheet.title)
            advisors = [canonical_adviser(text(row, 14 if legacy else 13)), canonical_adviser(text(row, 15 if legacy else 14))]
            examiners = [text(row, index) for index in ((20, 21, 22) if legacy else (19, 20, 21))]
            records.append({
                "name": name,
                "title": title,
                "year": year,
                "program": text(row, 1) if legacy else text(row, 30),
                "advisors": [item for item in advisors if item],
                "examiners": [item for item in examiners if item],
            })
    output.parent.mkdir(exist_ok=True)
    output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(records)} thesis records to {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    arguments = parser.parse_args()
    main(arguments.source, arguments.output)
